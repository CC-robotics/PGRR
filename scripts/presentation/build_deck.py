#!/usr/bin/env python3
# ruff: noqa: RUF001
"""Generate the 30-slide Chinese PGRR briefing from approved report data.

Slide specification and speaker notes can be checked without ``python-pptx``.
Actual PPTX generation fails with an actionable dependency message when that
package is unavailable.  No historical or live validation result is opened by
this script; it consumes only the stage-aware JSON emitted by the report asset
builder.
"""

from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import math
import shutil
import subprocess
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Any

PROJECT_ROOT = Path(__file__).resolve().parents[2]
PUBLIC_NAME = (
    "PGRR: Planning-Guided Failure-Triggered Recovery and Rejoin for Dynamic Social Navigation"
)
DEFAULT_REPORT_DATA = PROJECT_ROOT / "report/generated/report_data.json"
DEFAULT_OUTPUT = PROJECT_ROOT / "presentation/PGRR_report_zh.pptx"
DEFAULT_NOTES = PROJECT_ROOT / "presentation/speaker_notes_zh.md"

STATIC_ASSETS = {
    "paper/figures/system_architecture.pdf",
    "paper/figures/recovery_state_machine.pdf",
    "paper/figures/action_space_expert.pdf",
    "paper/figures/scenario_overview.pdf",
    "paper/figures/runtime_gazebo_doorway_bottleneck_medium.png",
}
RESULT_ASSETS = {
    "report/generated/result_outcomes.pdf",
    "report/generated/result_density.pdf",
    "report/generated/result_family.pdf",
    "report/generated/result_safety_efficiency.pdf",
    "report/generated/result_paired_effects.pdf",
    "report/generated/result_matched_run_evidence.pdf",
    "report/generated/result_matched_trajectory.pdf",
    "report/generated/result_matched_recovery_timeline.pdf",
}


class DeckBuildError(RuntimeError):
    """Raised when deck inputs or dependencies are incomplete."""


@dataclass(frozen=True)
class SlideSpec:
    number: int
    title: str
    takeaway: str
    bullets: tuple[str, ...]
    notes: str
    asset: str | None = None
    result_dependent: bool = False


def _percent(value: object) -> str:
    number = float(value)
    if not math.isfinite(number):
        raise DeckBuildError("non-finite value in approved report data")
    return f"{100.0 * number:.1f}%"


def _points(value: object) -> str:
    number = float(value)
    if not math.isfinite(number):
        raise DeckBuildError("non-finite paired effect in approved report data")
    return f"{100.0 * number:+.1f} 个百分点"


def load_report_data(path: Path, *, stage: str) -> dict[str, Any]:
    if stage not in {"pending", "validation", "test"}:
        raise DeckBuildError("stage must be pending, validation, or test")
    if not path.is_file():
        if stage == "pending":
            return {
                "schema_version": 1,
                "stage": "pending",
                "results_available": False,
                "author_alias": "Charles Chen",
            }
        raise DeckBuildError(
            f"stage {stage!r} requires report data generated from an approved result: {path}"
        )
    data = json.loads(path.read_text(encoding="utf-8"))
    if data.get("stage") != stage:
        raise DeckBuildError(
            f"report stage mismatch: requested {stage!r}, JSON declares {data.get('stage')!r}"
        )
    available = bool(data.get("results_available"))
    if stage == "pending" and available:
        raise DeckBuildError("pending deck must not contain numerical results")
    if stage != "pending" and not available:
        raise DeckBuildError(f"{stage} deck requires approved numerical report data")
    if data.get("author_alias") not in {None, "Charles Chen"}:
        raise DeckBuildError("presentation author must use the Charles Chen alias")
    if available:
        methods = {row.get("method") for row in data.get("method_summary", [])}
        expected = {"base", "standard", "heuristic", "bc_uniform", "pgrr"}
        if methods != expected:
            raise DeckBuildError("report data does not contain the complete five-method summary")
        if data.get("schema_version") != 3:
            raise DeckBuildError("result-bearing deck requires report data schema version 3")
        if data.get("public_name") != PUBLIC_NAME:
            raise DeckBuildError("report data does not use the exact public PGRR name")
        if data.get("benchmark_id") != "moderate_social_navigation_v6":
            raise DeckBuildError("result-bearing deck accepts only moderate-v6 evidence")
        if not str(data.get("statistics_sha256", "")).strip():
            raise DeckBuildError("report data omits paired-statistics provenance")
        paired = data.get("paired_comparisons")
        if not isinstance(paired, list) or len(paired) != 12:
            raise DeckBuildError("report data must contain all twelve paired binary effects")
        observed = {
            (str(row.get("comparator")), str(row.get("endpoint")))
            for row in paired
            if isinstance(row, dict)
        }
        expected_effects = {
            (comparator, endpoint)
            for comparator in ("base", "standard", "heuristic", "bc_uniform")
            for endpoint in ("goal_reached", "collision", "timeout")
        }
        if observed != expected_effects:
            raise DeckBuildError("report data paired effects are incomplete or duplicated")
        matched = data.get("matched_run_evidence")
        if not isinstance(matched, dict) or not bool(matched.get("available")):
            raise DeckBuildError("result-bearing deck requires matched Base--PGRR raw evidence")
        if matched.get("representation") != (
            "telemetry reconstruction; not a simulator camera screenshot"
        ):
            raise DeckBuildError("matched evidence must be labelled as telemetry reconstruction")
        if stage == "test":
            runtime = data.get("runtime_capture")
            if not isinstance(runtime, dict) or not bool(runtime.get("available")):
                raise DeckBuildError("locked-test deck requires verified Gazebo GUI evidence")
            for field in ("episode_id", "episode_outcome", "screenshot_sha256", "metadata_sha256"):
                if not str(runtime.get(field, "")).strip():
                    raise DeckBuildError(f"locked-test deck runtime provenance omits {field}")
            if len(str(runtime["screenshot_sha256"])) != 64:
                raise DeckBuildError("locked-test deck runtime screenshot SHA256 is invalid")
            if matched.get("artifact_type") != "matched_base_pgrr_test_media":
                raise DeckBuildError("locked-test deck requires schema-v2 matched test media")
            artifacts = matched.get("artifacts")
            expected_artifacts = {
                "trajectory": (
                    "moderate_matched_base_pgrr_trajectory.pdf",
                    PROJECT_ROOT / "report/generated/result_matched_trajectory.pdf",
                ),
                "recovery_timeline": (
                    "moderate_pgrr_recovery_timeline.pdf",
                    PROJECT_ROOT / "report/generated/result_matched_recovery_timeline.pdf",
                ),
            }
            if not isinstance(artifacts, dict) or set(artifacts) != set(expected_artifacts):
                raise DeckBuildError("locked-test deck requires trajectory and recovery timeline")
            for field, (filename, copied_asset) in expected_artifacts.items():
                declaration = artifacts[field]
                if not isinstance(declaration, dict) or declaration.get("filename") != filename:
                    raise DeckBuildError(f"locked-test deck {field} filename is not fixed")
                declared_sha = str(declaration.get("sha256", ""))
                if len(declared_sha) != 64 or not copied_asset.is_file():
                    raise DeckBuildError(f"locked-test deck {field} asset is incomplete")
                if _sha256(copied_asset) != declared_sha:
                    raise DeckBuildError(f"locked-test deck {field} asset disagrees with SHA256")
    return data


def _pending_result_bullets(page_purpose: str) -> tuple[str, ...]:
    return (
        "阶段标识：验证执行中，结果尚未锁定",
        f"本页计划展示：{page_purpose}",
        "未读取历史 64-episode、pilot、calibration 或运行中的验证目录",
        "批准结果通过 split、pair、完整性与 SHA256 检查后自动覆盖本页",
    )


def _summary_by_method(data: dict[str, Any]) -> dict[str, dict[str, Any]]:
    return {str(row["method"]): dict(row) for row in data["method_summary"]}


def _result_content(data: dict[str, Any]) -> dict[int, tuple[str, tuple[str, ...], str | None]]:
    if not bool(data.get("results_available")):
        return {
            21: (
                "只接受完整终止类别，不用安全替代完成",
                _pending_result_bullets("五种方法的到达、碰撞、超时与规划失败"),
                None,
            ),
            22: (
                "密度效应必须展示全部 low / medium / high 单元",
                _pending_result_bullets("三档密度的目标到达率与样本数"),
                None,
            ),
            23: (
                "八类场景全部公开，不能只挑有利案例",
                _pending_result_bullets("八个 family × 五种方法的完整目标到达矩阵"),
                None,
            ),
            24: (
                "配对效应与显著性必须来自锁定统计 JSON",
                _pending_result_bullets("PGRR 相对四个 baseline 的配对差、区间与 Holm 校正"),
                None,
            ),
            25: (
                "安全、效率与完成率必须联合解释",
                _pending_result_bullets("最小人距与成功 episode 导航时间"),
                None,
            ),
            26: (
                "matched 运行页等待 hash-linked raw/Parquet sidecar",
                _pending_result_bullets(
                    "预注册 doorway/medium/r00 的 Base--PGRR 空间轨迹与事件时间线；"
                    "遥测重建，不是 camera screenshot"
                ),
                None,
            ),
            27: (
                "恢复时序素材与轨迹素材必须来自同一 test pair",
                _pending_result_bullets(
                    "固定文件名的 PGRR 目标距离、failure score、恢复状态与触发时间线；"
                    "遥测重建，不是 camera screenshot"
                ),
                None,
            ),
        }

    summary = _summary_by_method(data)
    base = summary["base"]
    pgrr = summary["pgrr"]
    paired = {
        (str(row["comparator"]), str(row["endpoint"])): row for row in data["paired_comparisons"]
    }
    matched = data["matched_run_evidence"]
    matched_runs = matched["runs"]
    stage_label = "锁定 test" if data["stage"] == "test" else "validation 快照（非最终 test）"
    densities = {(row["density"], row["method"]): row for row in data["density_summary"]}
    density_lines = tuple(
        f"{density.title()}：DWB {_percent(densities[(density, 'base')]['goal_rate'])}；"
        f"PGRR {_percent(densities[(density, 'pgrr')]['goal_rate'])}"
        for density in ("low", "medium", "high")
    )
    result = {
        21: (
            f"{stage_label}：终止结果按四类完整报告",
            (
                f"配对条件：{data['condition_count']}；方法 episode：{data['episode_count']}",
                f"DWB 到达 / 碰撞 / 超时：{_percent(base['goal_rate'])} / "
                f"{_percent(base['collision_rate'])} / {_percent(base['timeout_rate'])}",
                f"PGRR 到达 / 碰撞 / 超时：{_percent(pgrr['goal_rate'])} / "
                f"{_percent(pgrr['collision_rate'])} / {_percent(pgrr['timeout_rate'])}",
                f"基础设施排除：{data['excluded_episode_count']}，未并入算法分母",
            ),
            "report/generated/result_outcomes.pdf",
        ),
        22: (
            "三档密度全部展示，不做事后子集选择",
            (*density_lines, "柱状图来自同一 approved results.parquet", f"阶段：{stage_label}"),
            "report/generated/result_density.pdf",
        ),
        23: (
            "聚合结果必须回到八类交互几何检查",
            (
                "每个单元显示目标到达率，覆盖全部预声明 family",
                "矩阵用于识别收益集中、退化和不可恢复结构",
                "场景差异为描述性分析，不自动构成显著性结论",
                f"阶段：{stage_label}",
            ),
            "report/generated/result_family.pdf",
        ),
        24: (
            "四个 baseline × 三个终止端点均给出配对差、区间、校正 p 与效应量",
            (
                *(
                    f"{label}：到达 {_points(paired[(method, 'goal_reached')]['difference'])}；"
                    f"碰撞 {_points(paired[(method, 'collision')]['difference'])}；"
                    f"超时 {_points(paired[(method, 'timeout')]['difference'])}"
                    for method, label in (
                        ("base", "DWB"),
                        ("standard", "Standard"),
                        ("heuristic", "Heuristic"),
                        ("bc_uniform", "Uniform BC"),
                    )
                ),
                "右图逐端点标注 95% CI、全局 Holm p_H 与 matched OR_H",
            ),
            "report/generated/result_paired_effects.pdf",
        ),
        25: (
            "安全—效率视图是补充，不替代终止结果",
            (
                "横轴：成功 episode 的中位导航时间",
                "纵轴：全部有效 episode 的中位最小人距",
                f"DWB 中位最小人距：{float(base['median_min_human_distance_m']):.2f} m",
                f"PGRR 中位最小人距：{float(pgrr['median_min_human_distance_m']):.2f} m",
                f"阶段：{stage_label}",
            ),
            "report/generated/result_safety_efficiency.pdf",
        ),
        26: (
            "固定 matched pair 公开 Base--PGRR 空间轨迹与真实终局",
            (
                f"选择规则：{matched['selection_rule']}",
                f"DWB 终局：{matched_runs['base']['outcome']}；"
                f"raw SHA {matched_runs['base']['raw_sha256'][:12]}",
                f"PGRR 终局：{matched_runs['pgrr']['outcome']}；"
                f"raw SHA {matched_runs['pgrr']['raw_sha256'][:12]}",
                f"pair：{matched['pair_id']}；scenario：{matched['scenario_id']}",
                "轨迹是 JSONL/Parquet 遥测重建，不是 camera screenshot",
            ),
            (
                "report/generated/result_matched_trajectory.pdf"
                if data["stage"] == "test"
                else "report/generated/result_matched_run_evidence.pdf"
            ),
        ),
        27: (
            "模型选择保留负面结果，不把离线指标包装成闭环收益",
            (
                "Uniform BC：相同观测、动作空间和规划 mask",
                "PGRR：validation 选择的 DAgger checkpoint",
                "margin weighting 若无优势则作为负面消融",
                "no-mask 只做离线反事实诊断，绝不执行",
                "闭环结论仍来自五方法相同 manifest",
            ),
            "paper/figures/action_space_expert.pdf",
        ),
    }
    if data["stage"] == "test":
        result[27] = (
            "恢复触发、状态切换和目标进展来自同一 PGRR raw stream",
            (
                f"pair：{matched['pair_id']}；PGRR episode：{matched_runs['pgrr']['episode_id']}",
                f"PGRR 终局：{matched_runs['pgrr']['outcome']}；"
                f"raw SHA {matched_runs['pgrr']['raw_sha256'][:12]}",
                "时间线包含记录的目标距离、failure score、恢复状态与触发事件",
                "时间线是 telemetry reconstruction，不是 camera screenshot",
                "n=1 描述性运行证据；总体结论仍来自完整 outcome 分解",
            ),
            "report/generated/result_matched_recovery_timeline.pdf",
        )
    return result


def build_slide_specs(stage: str, data: dict[str, Any]) -> tuple[SlideSpec, ...]:
    result = _result_content(data)
    runtime = data.get("runtime_capture")
    if isinstance(runtime, dict) and bool(runtime.get("available")):
        runtime_provenance = (
            f"episode ID：{runtime.get('episode_id')}；pixel SHA256："
            f"{str(runtime.get('screenshot_sha256', ''))[:12]}"
        )
    else:
        runtime_provenance = "截图、窗口、日志、episode、commit 与 SHA256 均由元数据绑定"
    stage_name = {
        "pending": "验证执行中｜数值待锁定",
        "validation": "Validation 快照｜非最终 Test",
        "test": "锁定 Test 结果",
    }[stage]

    specs = (
        SlideSpec(
            1,
            "PGRR：规划引导、失败触发的动态社会导航恢复与重接",
            PUBLIC_NAME,
            ("EI 会议项目汇报", stage_name, "Charles Chen"),
            "先明确今天汇报的范围：方法、真实运行证据和严格阶段化结果。当前阶段标签会出现在每一页。",
            "paper/figures/system_architecture.pdf",
        ),
        SlideSpec(
            2,
            "一句话结论",
            "学习模块不是新的底盘控制器，而是受规划约束的短时恢复决策层",
            (
                "正常状态继续使用 Nav2 DWB",
                "持续风险、冻结、振荡或死锁证据才触发恢复",
                "策略选择临时子目标或 WAIT / BACKUP / REPLAN / CONTINUE",
                "完成后恢复原始 PointGoal 并重新接回经典导航",
                "历史 v1 64/64（非 v6）：PGRR/Base 碰撞 0/24 vs 19/24、超时 "
                "16/24 vs 0/24、到达 8/24 vs 5/24；校正成功差异不显著",
            ),
            "用一句话把层级关系讲清楚，避免听众把 PGRR 误解为端到端速度策略。",
        ),
        SlideSpec(
            3,
            "为什么经典规划器仍会失败？",
            "局部可行不等于动态交互可恢复",
            (
                "对向会车：双方持续占据彼此最优局部轨迹",
                "门口竞争：短时最优控制造成冻结或相互抢占",
                "盲角突现：有限视野下接近速度快速变化",
                "临时封堵：局部规划器反复输出无效控制或振荡",
            ),
            "强调这些不是把 DWB 参数调坏制造的失败，而是动态社会交互下的局部决策问题。",
            "paper/figures/scenario_overview.pdf",
        ),
        SlideSpec(
            4,
            "失败不是单一碰撞标签",
            "系统区分碰撞风险、冻结、振荡和动态死锁",
            (
                "碰撞风险：前向硬防护或持续闭合趋势",
                "冻结：离目标仍远、位移不足且规划器请求运动",
                "振荡：角速度多次换向但目标进展不足",
                "死锁：持续阻塞且短窗口无法恢复有效进展",
                "最终结局仍独立记录为到达、碰撞、超时或规划失败",
            ),
            "这页要区分触发类型与 episode 终止类型，两者不能混为一个安全分数。",
        ),
        SlideSpec(
            5,
            "研究问题",
            "能否提高困难动态交互的恢复能力，同时不破坏正常规划稳定性？",
            (
                "介入应稀疏、可解释、可取消",
                "学习动作必须经过规划可行性约束",
                "训练可用特权监督，测试只能用机器人可观测信息",
                "效果必须在相同 episode manifest 上配对比较",
            ),
            "把研究问题落到四个可验证要求：稀疏介入、合法动作、信息隔离和配对评价。",
        ),
        SlideSpec(
            6,
            "贡献与边界",
            "贡献集中在失败触发、规划专家、恢复分布学习和受限重接",
            (
                "失败前触发的经典规划器恢复层",
                "特权短时域规划自动生成恢复示范",
                "Behavior Cloning + 两轮 DAgger 覆盖策略诱导状态",
                "规划 action mask + 独立安全监督 + 有界状态机",
                "PPO、学习 detector、第二 planner、Flatland、硬件与形式安全均非完成声明",
            ),
            "最后一条很重要：主动说明当前版本不依赖 PPO 或学习检测器来成立。",
        ),
        SlideSpec(
            7,
            "总体架构",
            "训练侧可用特权规划监督，部署侧严格闭合在 LiDAR 与 Nav2 状态上",
            (
                "观测构造 → 失败检测 → 滞回状态机",
                "候选动作 → action mask → 恢复策略",
                "Goal Mux 保存原目标并发送临时子目标",
                "DWB 执行动作，重接后恢复常态控制",
            ),
            "沿着图从左到右讲一遍，并指出虚线特权区域只在训练出现。",
            "paper/figures/system_architecture.pdf",
        ),
        SlideSpec(
            8,
            "正式观测与触发证据",
            "部署不使用行人真值、ID 或未来轨迹",
            (
                "最近 5 帧 × 180 beams LiDAR",
                "目标极坐标与前方 8 个局部路径点",
                "机器人速度与 DWB 当前输出",
                "10 步目标进展和角速度历史",
                "规划器状态与规则失败分数",
            ),
            "这里主动回答公平性问题：测试输入均可由机器人传感和导航栈产生。",
        ),
        SlideSpec(
            9,
            "25 个可解释恢复动作",
            "策略做高层选择，不直接输出底盘速度",
            (
                "21 个临时子目标：3 个半径 × 7 个相对方向",
                "4 个行为：WAIT、BACKUP、REPLAN、CONTINUE",
                "临时目标仍由 DWB 跟踪",
                "动作 ID 固定，训练、导出和 ROS 推理一致",
            ),
            "用图说明一个恢复动作最终仍经过经典局部规划，而不是绕过它。",
            "paper/figures/action_space_expert.pdf",
        ),
        SlideSpec(
            10,
            "Action Mask：先排除不可执行动作",
            "策略只能在物理和规划上可行的候选集合内选择",
            (
                "障碍内、地图外、局部不可达或不连通子目标被屏蔽",
                "明显进入动态占据区的子目标被屏蔽",
                "后方净空不足时禁止 BACKUP",
                "规划接口不可用时禁止 REPLAN",
                "masked invalid action rate 在部署中必须为 0",
            ),
            "强调 mask 是可行性约束，不是碰撞安全证明。",
        ),
        SlideSpec(
            11,
            "独立安全监督器",
            "学习决策永远不能覆盖紧急停止优先级",
            (
                "停止距离 = 制动距离 + 延迟距离 + margin",
                "窄前向区域保留即时硬防护",
                "偏轴风险需要时间一致的闭合证据",
                "旋转净空与平移足迹停止距离分离",
                "所有 emergency intervention 单独记录",
            ),
            "解释为什么要区分原地旋转与平移净空：安全侧墙不应永久锁死转向。",
        ),
        SlideSpec(
            12,
            "有界恢复状态机",
            "滞回、cooldown、动作保持、重接和最大尝试共同抑制抖动",
            (
                "NORMAL → PENDING → RECOVERY → REJOIN",
                "原目标在进入恢复时保存，临时目标结束后恢复",
                "REJOIN 只有重新产生原目标进展后才返回 NORMAL",
                "EMERGENCY STOP 可从任意执行状态抢占",
                "超时和连续失败进入明确终止状态",
            ),
            "按正常路径讲状态转换，再补充红色安全抢占和失败出口。",
            "paper/figures/recovery_state_machine.pdf",
        ),
        SlideSpec(
            13,
            "特权短时域规划专家",
            "用未来短窗口比较所有合法恢复动作，而不是人工逐帧标注",
            (
                "读取机器人、地图、行人位置速度和短时预测",
                "对每个合法候选进行差速运动学 rollout",
                "联合计算碰撞、进展、社会距离、重接、平滑与时间成本",
                "输出最优动作、25 维代价、mask 和 margin",
                "专家只用于训练与 Oracle 分析",
            ),
            "要明确 privileged 不等于测试作弊，因为专家输出的是训练标签。",
            "paper/figures/action_space_expert.pdf",
        ),
        SlideSpec(
            14,
            "Behavior Cloning 与两轮 DAgger",
            "DAgger 专门补充当前策略会访问、原始示范不足的恢复状态",
            (
                "LiDAR 1D CNN + 导航状态 MLP → 25 logits",
                "训练先完成 Uniform BC",
                "每轮在 train split 闭环运行并由专家重新标注",
                "聚合数据后只在固定 validation 上选择 checkpoint",
                "未选中的第二轮候选和 margin 负面结果继续保留",
            ),
            "不要只讲 top-1 accuracy；核心是策略访问分布和真实闭环选择。",
        ),
        SlideSpec(
            15,
            "训练—部署信息隔离",
            "强监督可以来自仿真真值，但部署接口必须保持可观测",
            (
                "privileged humans / future outcomes 只存在数据与专家侧",
                "导出模型不包含真值行人张量",
                "ROS 推理节点只接收标准化正式观测和 mask",
                "schema、checkpoint manifest 与接口测试共同检查泄漏",
            ),
            "这页回答审稿人常见问题：专家使用真值是否导致部署不可实现。",
            "paper/figures/system_architecture.pdf",
        ),
        SlideSpec(
            16,
            "真实 Arena/Gazebo 运行证据",
            "不是概念图：Jackal、动态行人、静态瓶颈和 Nav2 在同一 episode 实际运行",
            (
                "环境：Ubuntu 22.04 / ROS2 Humble / Arena Gazebo",
                "机器人与感知：Jackal / Nav2 DWB / 平面 LiDAR",
                "冻结 v5 演示：doorway_bottleneck / medium / validation",
                runtime_provenance,
                "它不是锁定 v6 统计回合的 camera frame，也不替代定量实验",
            ),
            "指出 Jackal、LiDAR 可见行人代理和门口几何；明确这是冻结场景运行证明，"
            "不能把遥测重建或该演示图冒充锁定 test 截图。",
            "paper/figures/runtime_gazebo_doorway_bottleneck_medium.png",
        ),
        SlideSpec(
            17,
            "八类场景 × 三档密度",
            "从正面对向到临时封堵，覆盖不同动态交互结构",
            (
                "head-on、doorway、crossing、blind corner",
                "group blocking、overtaking、opposite streams、temporary blockage",
                "每类 low / medium / high",
                "地图、起终点、行人路线和 seed 写入 manifest",
            ),
            "快速扫过八个小图，不逐个展开细节；强调固定模板和 seeded physical realization。",
            "paper/figures/scenario_overview.pdf",
        ),
        SlideSpec(
            18,
            "Split、规模与锁定规则",
            "validation 用于选择，test 只在代码、配置和 checkpoint 冻结后打开",
            (
                "moderate-v6：Train / validation / test 使用不相交 seed 和 scenario ID",
                "v5 Base validation 57/72，超过 75% ceiling，故 rejected 且 test 未打开",
                "Validation：72 条件 × 5 方法 = 360 method-episodes",
                "锁定 Test：120 条件 × 5 方法 = 600 method-episodes",
                "同一 pair 的场景、地图、seed 和行人配置跨方法一致",
                "test 结果不能反向调参",
            ),
            "明确 360 和 600 是总 method-episodes，不是每个方法的数量。",
        ),
        SlideSpec(
            19,
            "五种闭环对比方法",
            "从纯经典基线到训练分布聚合，逐级增加恢复能力",
            (
                "DWB：无 PGRR 恢复层",
                "Standard：导航栈标准恢复",
                "Heuristic：规则触发 + 手工动作",
                "Uniform BC：相同观测、动作和 mask 的行为克隆",
                "PGRR：选定 DAgger checkpoint + 有界恢复闭环",
            ),
            "强调所有方法复用同一 DWB 和同一物理条件，避免基础规划器差异干扰比较。",
        ),
        SlideSpec(
            20,
            "指标与统计协议",
            "一个方法不能靠永远 WAIT 获得虚假安全优势",
            (
                "终止：到达、碰撞、超时、规划失败",
                "效率：SPL、路径长度、导航时间",
                "安全：最小人距、个人空间侵入、不舒适时间、紧急停止",
                "恢复：触发、重接成功、持续时间、介入比例",
                "配对 McNemar / Wilcoxon + bootstrap 95% CI + 全局 Holm",
            ),
            "这页为后面的结果解释定规则：显著性、效果量和失败类别都要一起看。",
        ),
        SlideSpec(
            21,
            "主结果：完整终止类别",
            result[21][0],
            result[21][1],
            "先确认阶段标签，再报告所有终止类别。禁止只展示成功率或只展示碰撞率。",
            result[21][2],
            True,
        ),
        SlideSpec(
            22,
            "密度分层",
            result[22][0],
            result[22][1],
            "按 low、medium、high 顺序解释趋势；不要把描述性差异说成显著交互效应。",
            result[22][2],
            True,
        ),
        SlideSpec(
            23,
            "八类交互族结果",
            result[23][0],
            result[23][1],
            "完整矩阵用于定位方法在哪些几何交互中受益或退化，不能只截取表现好的三类。",
            result[23][2],
            True,
        ),
        SlideSpec(
            24,
            "PGRR 对 Baseline 的配对效应",
            result[24][0],
            result[24][1],
            "正的目标差有利；碰撞和超时则负值有利。逐项读取区间、全局 Holm pH 与 ORH。",
            result[24][2],
            True,
        ),
        SlideSpec(
            25,
            "安全—效率视图",
            result[25][0],
            result[25][1],
            "先讲坐标含义，再强调它不能把失败 episode 从完成率中删除。",
            result[25][2],
            True,
        ),
        SlideSpec(
            26,
            "Matched Base--PGRR 真实运行对比",
            result[26][0],
            result[26][1],
            "先核对 pair 与 raw SHA，再读轨迹和事件线；这是遥测重建，不是相机截图。",
            result[26][2],
            True,
        ),
        SlideSpec(
            27,
            (
                "PGRR 恢复时序（同一 Matched Pair）"
                if stage in {"pending", "test"}
                else "训练与消融"
            ),
            result[27][0],
            result[27][1],
            (
                "核对同一 pair、PGRR episode 和 raw SHA 后再读触发与状态线；"
                "这是遥测重建，不是相机画面。"
                if stage in {"pending", "test"}
                else "把 checkpoint 选择和负面结果讲清楚：未带来验证提升的模块不进入贡献结论。"
            ),
            result[27][2],
            True,
        ),
        SlideSpec(
            28,
            "失败案例与局限",
            "降低碰撞但增加 timeout 仍然是失败转移，必须公开",
            (
                "重点检查长期 WAIT、持续 BACKUP、左右切换和重复恢复",
                "规则触发可能误报或漏报，action mask 可能过于保守",
                "二维 LiDAR、已知地图、离散动作和仿真人群限制外推",
                "历史 v1 64/64（非 v6）显示碰撞下降伴随 timeout 上升；"
                "到达 8/24 vs 5/24 的校正比较不显著",
                "PPO、学习 detector、第二 planner、Flatland、硬件和形式安全均未完成",
            ),
            "结合代表性轨迹解释根因；不要用截图代替总体失败统计。",
            "paper/figures/recovery_state_machine.pdf",
        ),
        SlideSpec(
            29,
            "复现与工程交付",
            "代码、结果、图表、论文和汇报共享同一证据链",
            (
                "锁定 episode manifest、run manifest 和 results.parquet",
                "保存项目 commit、Arena commit、checkpoint 与场景 SHA256",
                "图表和 TeX 表格全部自动生成",
                "发布门禁：test 技术报告固定 32 页；PPTX/PDF 固定 30 页",
                "发布前执行测试、字体、关系、占位符和隐私审计",
            ),
            "展示一键命令，并说明任何完整性检查失败都会阻止生成“final”文档。",
        ),
        SlideSpec(
            30,
            "结论与 Q&A",
            "PGRR 的核心不是替代经典规划，而是让失败恢复可学习、可约束、可解释",
            (
                "经典规划器保持常态控制",
                "特权规划专家降低人工恢复标注成本",
                "DAgger 覆盖策略诱导的困难恢复状态",
                "规划 mask 与有界状态机限制学习策略作用域",
                f"当前汇报阶段：{stage_name}",
            ),
            "最后再次说明阶段：如果还是 pending，只总结已验证系统和协议，"
            "不口头补入未经锁定的数字。",
            "paper/figures/system_architecture.pdf",
        ),
    )
    if len(specs) != 30 or tuple(spec.number for spec in specs) != tuple(range(1, 31)):
        raise DeckBuildError("deck specification must contain exactly slides 1..30")
    return specs


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def validate_runtime_capture() -> None:
    screenshot = PROJECT_ROOT / "paper/figures/runtime_gazebo_doorway_bottleneck_medium.png"
    source = PROJECT_ROOT / "outputs/figures/runtime/gazebo_doorway_bottleneck_medium.png"
    metadata_path = (
        PROJECT_ROOT / "outputs/figures/runtime/gazebo_doorway_bottleneck_medium.metadata.json"
    )
    window_path = (
        PROJECT_ROOT / "outputs/figures/runtime/gazebo_doorway_bottleneck_medium.window.json"
    )
    for path in (screenshot, source, metadata_path, window_path):
        if not path.is_file():
            raise DeckBuildError(f"verified Gazebo capture artifact is missing: {path.name}")
    metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
    if metadata.get("artifact_type") != "real_arena_gazebo_gui_screenshot":
        raise DeckBuildError("runtime capture is not declared as a real Gazebo GUI screenshot")
    declared = str(metadata.get("artifacts", {}).get("screenshot_sha256", ""))
    if len(declared) != 64 or _sha256(source) != declared or _sha256(screenshot) != declared:
        raise DeckBuildError("runtime capture pixels disagree with their declared SHA256")
    launch = metadata.get("launch_profile", {})
    if (
        launch.get("simulator") != "gazebo"
        or launch.get("robot") != "jackal"
        or launch.get("local_planner") != "dwb"
    ):
        raise DeckBuildError("runtime capture must document Gazebo/Jackal/Nav2 DWB")
    scenario = metadata.get("scenario", {})
    if scenario.get("split") != "validation" or "moderate_v5" not in str(
        scenario.get("scenario_id", "")
    ):
        raise DeckBuildError("runtime capture must remain the audited frozen-v5 demo")
    for field in ("episode_id", "episode_outcome", "git_commit", "arena_commit"):
        if not str(metadata.get(field, "")).strip():
            raise DeckBuildError(f"runtime capture provenance omits {field}")
    if len(str(metadata["git_commit"])) != 40 or len(str(metadata["arena_commit"])) != 40:
        raise DeckBuildError("runtime capture provenance contains an invalid revision")
    if "same Arena/Nav2 episode" not in str(metadata.get("provenance_note", "")):
        raise DeckBuildError("runtime capture pixels are not episode-bound")
    window = json.loads(window_path.read_text(encoding="utf-8"))
    if window.get("selected_window", {}).get("title") != "Gazebo":
        raise DeckBuildError("runtime capture window metadata is invalid")


def validate_assets(specs: tuple[SlideSpec, ...], *, stage: str) -> None:
    validate_runtime_capture()
    allowed = STATIC_ASSETS | (RESULT_ASSETS if stage != "pending" else set())
    for spec in specs:
        if spec.asset is None:
            continue
        if spec.asset not in allowed:
            raise DeckBuildError(f"slide {spec.number} uses an unapproved asset: {spec.asset}")
        path = (PROJECT_ROOT / spec.asset).resolve()
        try:
            path.relative_to(PROJECT_ROOT)
        except ValueError as exc:
            raise DeckBuildError(f"slide asset escapes the repository: {spec.asset}") from exc
        if not path.is_file():
            raise DeckBuildError(f"slide {spec.number} asset is missing: {spec.asset}")


def render_notes(specs: tuple[SlideSpec, ...], *, stage: str) -> str:
    lines = [
        "# PGRR 汇报逐页讲稿",
        "",
        f"阶段：`{stage}`。讲稿与 PPT 由同一 30 页 slide specification 生成。",
        "",
    ]
    for spec in specs:
        lines.extend(
            (
                f"## {spec.number:02d}. {spec.title}",
                "",
                f"核心句：{spec.takeaway}",
                "",
                spec.notes,
                "",
                "讲述要点：",
                "",
            )
        )
        lines.extend(f"- {bullet}" for bullet in spec.bullets)
        lines.append("")
    return "\n".join(lines)


def _require_pptx() -> dict[str, Any]:
    if importlib.util.find_spec("pptx") is None:
        raise DeckBuildError(
            "python-pptx is required to generate PPTX but is not installed in ramp-offline. "
            "After the running evaluation finishes, add/install python-pptx in ramp-offline "
            "and rerun 'make presentation'. No ROS/Arena environment should be modified."
        )
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt

    return {
        "Presentation": Presentation,
        "RGBColor": RGBColor,
        "MSO_SHAPE": MSO_SHAPE,
        "PP_ALIGN": PP_ALIGN,
        "Inches": Inches,
        "Pt": Pt,
    }


def _raster_asset(asset: Path, generated_dir: Path) -> Path:
    if asset.suffix.lower() in {".png", ".jpg", ".jpeg"}:
        return asset
    if asset.suffix.lower() != ".pdf":
        raise DeckBuildError(f"unsupported slide asset format: {asset}")
    executable = shutil.which("pdftoppm")
    if executable is None:
        raise DeckBuildError("pdftoppm is required to rasterize vector figures for PPTX")
    generated_dir.mkdir(parents=True, exist_ok=True)
    output_stem = generated_dir / asset.stem
    output = output_stem.with_suffix(".png")
    subprocess.run(
        [executable, "-png", "-singlefile", "-r", "210", str(asset), str(output_stem)],
        check=True,
        capture_output=True,
        text=True,
    )
    if not output.is_file() or output.stat().st_size == 0:
        raise DeckBuildError(f"failed to rasterize slide asset: {asset}")
    return output


def _add_textbox(
    slide: Any,
    api: dict[str, Any],
    x: float,
    y: float,
    w: float,
    h: float,
    text: str,
    *,
    size: float,
    bold: bool = False,
    color: tuple[int, int, int] = (27, 31, 35),
    align: Any | None = None,
) -> Any:
    shape = slide.shapes.add_textbox(
        api["Inches"](x), api["Inches"](y), api["Inches"](w), api["Inches"](h)
    )
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.font.name = "Noto Sans CJK SC"
    paragraph.font.size = api["Pt"](size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = api["RGBColor"](*color)
    if align is not None:
        paragraph.alignment = align
    return shape


def _add_picture_fit(
    slide: Any, api: dict[str, Any], path: Path, x: float, y: float, w: float, h: float
) -> None:
    from PIL import Image

    with Image.open(path) as image:
        image_w, image_h = image.size
    box_ratio = w / h
    image_ratio = image_w / image_h
    if image_ratio >= box_ratio:
        draw_w = w
        draw_h = w / image_ratio
        draw_x = x
        draw_y = y + (h - draw_h) / 2.0
    else:
        draw_h = h
        draw_w = h * image_ratio
        draw_x = x + (w - draw_w) / 2.0
        draw_y = y
    slide.shapes.add_picture(
        str(path),
        api["Inches"](draw_x),
        api["Inches"](draw_y),
        width=api["Inches"](draw_w),
        height=api["Inches"](draw_h),
    )


def _add_bullets(
    slide: Any,
    api: dict[str, Any],
    bullets: tuple[str, ...],
    *,
    x: float,
    y: float,
    w: float,
    h: float,
    size: float,
) -> None:
    shape = slide.shapes.add_textbox(
        api["Inches"](x), api["Inches"](y), api["Inches"](w), api["Inches"](h)
    )
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    for index, bullet in enumerate(bullets):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = f"• {bullet}"
        paragraph.font.name = "Noto Sans CJK SC"
        paragraph.font.size = api["Pt"](size)
        paragraph.font.color.rgb = api["RGBColor"](27, 31, 35)
        paragraph.space_after = api["Pt"](8)


def generate_pptx(specs: tuple[SlideSpec, ...], *, stage: str, output: Path) -> None:
    api = _require_pptx()
    presentation = api["Presentation"]()
    presentation.slide_width = api["Inches"](13.333)
    presentation.slide_height = api["Inches"](7.5)
    presentation.core_properties.title = PUBLIC_NAME
    presentation.core_properties.subject = f"stage={stage}; 30-slide reproducible briefing"
    presentation.core_properties.author = "Charles Chen"
    presentation.core_properties.last_modified_by = "Charles Chen"
    generated_dir = PROJECT_ROOT / "presentation/generated"
    blank_layout = presentation.slide_layouts[6]
    stage_text = {
        "pending": "验证执行中｜结果待锁定",
        "validation": "Validation 快照｜非最终 Test",
        "test": "锁定 Test 结果",
    }[stage]

    for spec in specs:
        slide = presentation.slides.add_slide(blank_layout)
        background = slide.background.fill
        background.solid()
        background.fore_color.rgb = api["RGBColor"](255, 255, 255)

        if spec.number == 1:
            _add_textbox(slide, api, 0.8, 0.65, 11.7, 1.15, spec.title, size=29, bold=True)
            _add_textbox(
                slide, api, 0.8, 1.82, 11.7, 0.70, spec.takeaway, size=19, color=(53, 110, 159)
            )
            if spec.asset:
                raster = _raster_asset(PROJECT_ROOT / spec.asset, generated_dir)
                _add_picture_fit(slide, api, raster, 0.85, 2.55, 8.3, 3.95)
            _add_bullets(slide, api, spec.bullets, x=9.35, y=3.0, w=3.1, h=2.5, size=17)
        else:
            _add_textbox(slide, api, 0.65, 0.32, 11.5, 0.62, spec.title, size=25, bold=True)
            accent = slide.shapes.add_shape(
                api["MSO_SHAPE"].RECTANGLE,
                api["Inches"](0.65),
                api["Inches"](1.02),
                api["Inches"](12.0),
                api["Inches"](0.03),
            )
            accent.fill.solid()
            accent.fill.fore_color.rgb = api["RGBColor"](213, 94, 0)
            accent.line.fill.background()
            _add_textbox(
                slide,
                api,
                0.68,
                1.10,
                11.8,
                0.58,
                spec.takeaway,
                size=17.5,
                bold=True,
                color=(53, 110, 159),
            )
            if spec.asset:
                raster = _raster_asset(PROJECT_ROOT / spec.asset, generated_dir)
                if spec.number == 24:
                    _add_picture_fit(slide, api, raster, 0.72, 1.70, 11.9, 5.18)
                else:
                    _add_bullets(slide, api, spec.bullets, x=0.75, y=1.85, w=5.0, h=4.8, size=16.5)
                    _add_picture_fit(slide, api, raster, 5.9, 1.75, 6.7, 4.95)
            else:
                _add_bullets(slide, api, spec.bullets, x=1.0, y=1.88, w=11.2, h=4.8, size=19)

        badge_color = (213, 94, 0) if stage == "pending" else (53, 110, 159)
        _add_textbox(slide, api, 0.65, 7.05, 7.6, 0.25, stage_text, size=9.5, color=badge_color)
        _add_textbox(
            slide,
            api,
            10.3,
            7.05,
            2.35,
            0.25,
            f"PGRR  |  {spec.number:02d}/30",
            size=9.5,
            color=(89, 99, 107),
            align=api["PP_ALIGN"].RIGHT,
        )

    output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output)
    if not output.is_file() or output.stat().st_size == 0:
        raise DeckBuildError(f"PPTX was not created: {output}")


def _parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--stage", choices=("pending", "validation", "test"), default="pending")
    parser.add_argument("--report-data", type=Path, default=DEFAULT_REPORT_DATA)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--notes", type=Path, default=DEFAULT_NOTES)
    parser.add_argument("--check-only", action="store_true")
    parser.add_argument("--notes-only", action="store_true")
    return parser


def main() -> int:
    args = _parser().parse_args()
    try:
        data = load_report_data(args.report_data, stage=args.stage)
        specs = build_slide_specs(args.stage, data)
        validate_assets(specs, stage=args.stage)
        if args.check_only:
            print(f"Presentation specification PASS: 30 slides, stage={args.stage}")
            return 0
        notes_payload = render_notes(specs, stage=args.stage)
        if args.notes_only:
            args.notes.parent.mkdir(parents=True, exist_ok=True)
            args.notes.write_text(notes_payload, encoding="utf-8")
            print(f"Speaker notes PASS: {args.notes}")
            return 0
        generate_pptx(specs, stage=args.stage, output=args.output)
        args.notes.parent.mkdir(parents=True, exist_ok=True)
        args.notes.write_text(notes_payload, encoding="utf-8")
    except (DeckBuildError, OSError, ValueError, subprocess.CalledProcessError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 2
    print(f"Presentation PASS: {args.output} (30 slides, stage={args.stage})")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
